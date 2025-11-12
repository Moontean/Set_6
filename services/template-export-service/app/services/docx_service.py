from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO

def generate_docx(cv_data: dict) -> bytes:
    """Generate DOCX from CV data"""
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    # Name (Title)
    name_paragraph = doc.add_paragraph()
    name_run = name_paragraph.add_run(cv_data.get('fullName', 'N/A'))
    name_run.font.size = Pt(24)
    name_run.font.bold = True
    name_run.font.color.rgb = RGBColor(44, 62, 80)
    name_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Contact Information
    contact_info = []
    if cv_data.get('email'):
        contact_info.append(cv_data['email'])
    if cv_data.get('phone'):
        contact_info.append(cv_data['phone'])
    if cv_data.get('address'):
        contact_info.append(cv_data['address'])
    
    if contact_info:
        contact_paragraph = doc.add_paragraph(' | '.join(contact_info))
        contact_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()  # Add space

    # Professional Summary
    if cv_data.get('summary'):
        heading = doc.add_heading('Professional Summary', level=2)
        heading.runs[0].font.color.rgb = RGBColor(52, 73, 94)
        doc.add_paragraph(cv_data['summary'])

    # Experience
    if cv_data.get('experience'):
        heading = doc.add_heading('Experience', level=2)
        heading.runs[0].font.color.rgb = RGBColor(52, 73, 94)
        
        for exp in cv_data['experience']:
            # Position and Company
            position_para = doc.add_paragraph()
            position_run = position_para.add_run(
                f"{exp.get('position', 'N/A')} at {exp.get('company', 'N/A')}"
            )
            position_run.bold = True
            
            # Dates
            dates = f"{exp.get('startDate', 'N/A')} - {exp.get('endDate', 'Present')}"
            doc.add_paragraph(dates)
            
            # Description
            if exp.get('description'):
                doc.add_paragraph(exp['description'])
            
            doc.add_paragraph()  # Add space

    # Education
    if cv_data.get('education'):
        heading = doc.add_heading('Education', level=2)
        heading.runs[0].font.color.rgb = RGBColor(52, 73, 94)
        
        for edu in cv_data['education']:
            # Degree and Field
            degree_para = doc.add_paragraph()
            degree_run = degree_para.add_run(
                f"{edu.get('degree', 'N/A')} in {edu.get('field', 'N/A')}"
            )
            degree_run.bold = True
            
            # Institution
            doc.add_paragraph(edu.get('institution', 'N/A'))
            
            # Dates
            dates = f"{edu.get('startDate', 'N/A')} - {edu.get('endDate', 'Present')}"
            doc.add_paragraph(dates)
            doc.add_paragraph()  # Add space

    # Skills
    if cv_data.get('skills'):
        heading = doc.add_heading('Skills', level=2)
        heading.runs[0].font.color.rgb = RGBColor(52, 73, 94)
        skills_text = ', '.join(cv_data['skills'])
        doc.add_paragraph(skills_text)

    # Languages
    if cv_data.get('languages'):
        heading = doc.add_heading('Languages', level=2)
        heading.runs[0].font.color.rgb = RGBColor(52, 73, 94)
        languages_text = ', '.join(cv_data['languages'])
        doc.add_paragraph(languages_text)

    # Save to BytesIO
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
